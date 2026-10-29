import ctypes
import io
import json
import re
import threading
import tkinter as tk
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from tkinter import font as tkfont
from tkinter import filedialog, messagebox, ttk
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0 Safari/537.36"
)
MAX_IMAGE_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class SharedImage:
    placeholder: str
    data: bytes


@dataclass(frozen=True)
class PastedTextFormat:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    color: str | None = None
    background: str | None = None
    font_name: str | None = None
    font_size: float | None = None


@dataclass(frozen=True)
class PastedTextRun:
    text: str
    format: PastedTextFormat = PastedTextFormat()


@dataclass(frozen=True)
class PastedTextParagraph:
    runs: tuple[PastedTextRun, ...]


def _css_color(value):
    value = value.strip().lower()
    if value in {"", "transparent", "initial", "inherit", "none"}:
        return None
    named_colors = {
        "black": "000000",
        "white": "FFFFFF",
        "red": "FF0000",
        "green": "008000",
        "blue": "0000FF",
        "yellow": "FFFF00",
        "gray": "808080",
        "grey": "808080",
        "gray25": "D9D9D9",
        "grey25": "D9D9D9",
        "lightgray": "D3D3D3",
        "lightgrey": "D3D3D3",
        "silver": "C0C0C0",
        "maroon": "800000",
        "purple": "800080",
        "fuchsia": "FF00FF",
        "lime": "00FF00",
        "olive": "808000",
        "navy": "000080",
        "teal": "008080",
        "aqua": "00FFFF",
        "orange": "FFA500",
    }
    if value in named_colors:
        return named_colors[value]
    if value.startswith("#"):
        digits = value[1:]
        if re.fullmatch(r"[0-9a-f]{3}", digits):
            return "".join(character * 2 for character in digits).upper()
        if re.fullmatch(r"[0-9a-f]{6}", digits):
            return digits.upper()
    match = re.fullmatch(
        r"rgba?\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})"
        r"(?:\s*,\s*[\d.]+)?\s*\)",
        value,
    )
    if match:
        channels = tuple(int(channel) for channel in match.groups())
        if all(channel <= 255 for channel in channels):
            return "".join(f"{channel:02X}" for channel in channels)
    return None


def _style_declarations(value):
    styles = {}
    for declaration in value.split(";"):
        if ":" not in declaration:
            continue
        name, content = declaration.split(":", 1)
        styles[name.strip().lower()] = content.strip()
    return styles


class _PastedHtmlParser(HTMLParser):
    BLOCK_TAGS = {
        "address",
        "article",
        "blockquote",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "p",
        "pre",
    }

    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.paragraphs = [[]]
        self.style_stack = [PastedTextFormat()]
        self.open_tags = []
        self.ignored_tags = []
        self.rules = self._css_rules(html)
        self.feed(html)
        self.close()
        self._trim_paragraph()
        while len(self.paragraphs) > 1 and not self.paragraphs[-1]:
            self.paragraphs.pop()

    @staticmethod
    def _css_rules(html):
        rules = {}
        for block in re.findall(
            r"<style\b[^>]*>(.*?)</style\s*>",
            html,
            flags=re.IGNORECASE | re.DOTALL,
        ):
            for selectors, declarations in re.findall(
                r"([^{}]+)\{([^{}]*)\}",
                block,
            ):
                for selector in selectors.split(","):
                    selector = selector.strip()
                    if selector:
                        rules[selector] = _style_declarations(declarations)
        return rules

    def _trim_paragraph(self):
        if not self.paragraphs or not self.paragraphs[-1]:
            return
        runs = self.paragraphs[-1]
        first = runs[0]
        trimmed = first.text.lstrip()
        if trimmed != first.text:
            if trimmed:
                runs[0] = PastedTextRun(trimmed, first.format)
            else:
                runs.pop(0)
        if runs:
            last = runs[-1]
            trimmed = last.text.rstrip()
            if trimmed:
                runs[-1] = PastedTextRun(trimmed, last.format)
            else:
                runs.pop()

    def _next_paragraph(self):
        self._trim_paragraph()
        if self.paragraphs[-1]:
            self.paragraphs.append([])

    def _current_style(self, tag, attributes):
        current = self.style_stack[-1]
        declarations = {}
        classes = dict(attributes).get("class", "").split()
        for class_name in classes:
            declarations.update(
                self.rules.get(f".{class_name}", {})
            )
            declarations.update(
                self.rules.get(f"{tag}.{class_name}", {})
            )
        declarations.update(
            self.rules.get(tag, {})
        )
        attributes = dict(attributes)
        declarations.update(_style_declarations(attributes.get("style", "")))

        bold = current.bold or tag in {"b", "strong"}
        italic = current.italic or tag in {"i", "em"}
        underline = current.underline or tag in {"u", "ins"}
        if declarations.get("font-weight", "").lower() in {
            "bold",
            "bolder",
            "600",
            "700",
            "800",
            "900",
        }:
            bold = True
        if declarations.get("font-style", "").lower() in {
            "italic",
            "oblique",
        }:
            italic = True
        if "underline" in declarations.get("text-decoration", "").lower():
            underline = True

        color = current.color
        background = current.background
        font_name = current.font_name
        font_size = current.font_size
        if "color" in declarations:
            color = _css_color(declarations["color"]) or color
        background_value = declarations.get(
            "background-color",
            declarations.get("background", ""),
        )
        if not background_value and declarations.get("mso-highlight"):
            background_value = declarations["mso-highlight"]
        if background_value:
            background = _css_color(background_value) or background
        if attributes.get("color"):
            color = _css_color(attributes["color"]) or color
        if attributes.get("face"):
            font_name = attributes["face"].split(",")[0].strip(" '\"")
        if declarations.get("font-family"):
            font_name = (
                declarations["font-family"]
                .split(",")[0]
                .strip(" '\"")
            )
        size = declarations.get("font-size", "")
        size_match = re.fullmatch(r"([\d.]+)\s*(pt|px)?", size.lower())
        if size_match:
            font_size = float(size_match.group(1))
            if size_match.group(2) == "px":
                font_size *= 0.75
        elif attributes.get("size", "").isdigit():
            font_size = (int(attributes["size"]) + 5) * 2
        return PastedTextFormat(
            bold=bold,
            italic=italic,
            underline=underline,
            color=color,
            background=background,
            font_name=font_name,
            font_size=font_size,
        )

    def _append_text(self, text):
        text = re.sub(r"[ \t\r\n\f]+", " ", text)
        if not text:
            return
        runs = self.paragraphs[-1]
        current_format = self.style_stack[-1]
        if runs and runs[-1].format == current_format:
            previous = runs[-1]
            if previous.text.endswith(" ") and text.startswith(" "):
                text = text[1:]
            if text:
                runs[-1] = PastedTextRun(previous.text + text, current_format)
        elif text:
            runs.append(PastedTextRun(text, current_format))

    def handle_starttag(self, tag, attrs):
        if tag in {"head", "style", "script", "title"}:
            self.ignored_tags.append(tag)
            return
        if self.ignored_tags:
            return
        if tag == "br":
            self._next_paragraph()
            return
        if tag in {"meta", "link", "img", "hr", "input"}:
            return
        if tag in self.BLOCK_TAGS:
            self._next_paragraph()
        self.style_stack.append(self._current_style(tag, attrs))
        self.open_tags.append(tag)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.ignored_tags:
            if tag == self.ignored_tags[-1]:
                self.ignored_tags.pop()
            return
        if tag in self.BLOCK_TAGS:
            self._next_paragraph()
        if tag in self.open_tags:
            matching_tag = len(self.open_tags) - 1 - self.open_tags[::-1].index(tag)
            del self.open_tags[matching_tag:]
            del self.style_stack[matching_tag + 1 :]

    def handle_data(self, data):
        if not self.ignored_tags:
            self._append_text(data)


def parse_pasted_html(html):
    style_blocks = re.findall(
        r"<style\b[^>]*>.*?</style\s*>",
        html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    html_without_styles = re.sub(
        r"<style\b[^>]*>.*?</style\s*>",
        "",
        html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    parser = _PastedHtmlParser(
        "\n".join(style_blocks) + html_without_styles
    )
    return tuple(
        PastedTextParagraph(tuple(paragraph))
        for paragraph in parser.paragraphs
        if paragraph
    )


def _clipboard_html():
    if not hasattr(ctypes, "windll"):
        return None
    from ctypes import wintypes

    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    clipboard_format = user32.RegisterClipboardFormatW("HTML Format")
    if not clipboard_format:
        return None
    user32.OpenClipboard.argtypes = [wintypes.HWND]
    user32.OpenClipboard.restype = wintypes.BOOL
    user32.GetClipboardData.argtypes = [wintypes.UINT]
    user32.GetClipboardData.restype = wintypes.HANDLE
    kernel32.GlobalLock.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalSize.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalSize.restype = ctypes.c_size_t
    kernel32.GlobalUnlock.argtypes = [wintypes.HGLOBAL]
    kernel32.GlobalUnlock.restype = wintypes.BOOL
    if not user32.OpenClipboard(None):
        return None
    try:
        handle = user32.GetClipboardData(clipboard_format)
        if not handle:
            return None
        pointer = kernel32.GlobalLock(handle)
        if not pointer:
            return None
        try:
            raw_html = ctypes.string_at(pointer, kernel32.GlobalSize(handle))
        finally:
            kernel32.GlobalUnlock(handle)
    finally:
        user32.CloseClipboard()

    header = raw_html[:4096]
    match = re.search(rb"StartFragment:(\d+)", header)
    end_match = re.search(rb"EndFragment:(\d+)", header)
    if match and end_match:
        start = int(match.group(1))
        end = int(end_match.group(1))
        if 0 <= start <= end <= len(raw_html):
            raw_html = raw_html[start:end]
    else:
        start_marker = raw_html.find(b"<!--StartFragment-->")
        end_marker = raw_html.find(b"<!--EndFragment-->")
        if start_marker >= 0 and end_marker >= start_marker:
            raw_html = raw_html[
                start_marker + len(b"<!--StartFragment-->") : end_marker
            ]
    return raw_html.decode("utf-8-sig", errors="replace")


def validate_share_url(url):
    parsed = urlparse(url.strip())
    allowed_hosts = {"chatgpt.com", "www.chatgpt.com", "chat.openai.com"}
    if (
        parsed.scheme != "https"
        or parsed.hostname not in allowed_hosts
        or not (
            parsed.path.startswith("/share/")
            or parsed.path.startswith("/s/")
        )
    ):
        raise ValueError(
            "Enter a public ChatGPT share link beginning with "
            "https://chatgpt.com/share/ or https://chatgpt.com/s/."
        )
    return parsed.geturl()


def _content_text(content):
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        return "\n".join(
            text for item in content if (text := _content_text(item))
        ).strip()
    if isinstance(content, dict):
        if isinstance(content.get("text"), str):
            return content["text"].strip()
        if "parts" in content:
            return _content_text(content["parts"])
    return ""


def _with_citation_titles(text, references):
    for reference in references:
        if not isinstance(reference, dict):
            continue
        marker = reference.get("matched_text")
        alt = reference.get("alt")
        if not isinstance(marker, str) or not marker.startswith("\ue200cite"):
            continue
        if not isinstance(alt, str):
            text = text.replace(marker, "")
            continue
        titles = re.findall(r"\[([^\]]+)\]\([^)]+\)", alt)
        if not titles:
            text = text.replace(marker, "")
            continue

        def replace_marker(_match):
            start = text.rfind("\n", 0, _match.start()) + 1
            end = text.find("\n", _match.end())
            paragraph = text[start:] if end == -1 else text[start:end]
            return "" if any(
                title.casefold() in paragraph.casefold()
                for title in titles
            ) else ", ".join(titles)

        text = re.sub(re.escape(marker), replace_marker, text)
    return text


def _assistant_message_records(value):
    records = []

    def visit(item):
        if isinstance(item, dict):
            author = item.get("author")
            role = author.get("role") if isinstance(author, dict) else item.get("role")
            content = item.get("content")
            if role == "assistant" and content is not None:
                text = _content_text(content)
                if text:
                    metadata = item.get("metadata")
                    references = (
                        metadata.get("content_references", [])
                        if isinstance(metadata, dict)
                        else []
                    )
                    if isinstance(references, list):
                        text = _with_citation_titles(text, references)
                    records.append((text, metadata))
            for child in item.values():
                visit(child)
        elif isinstance(item, list):
            for child in item:
                visit(child)

    visit(value)
    return records


def _assistant_messages(value):
    return [text for text, _metadata in _assistant_message_records(value)]


def _latest_record_from_json_scripts(soup):
    candidates = []
    for script in soup.find_all("script"):
        raw = script.string or script.get_text()
        raw = raw.strip()
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            continue
        records = _assistant_message_records(data)
        if records:
            candidates.append(records)

    if not candidates:
        return None
    return max(candidates, key=len)[-1]


def _latest_record_from_react_router_stream(soup):
    candidates = []
    stream_pattern = re.compile(
        r"streamController\.enqueue\((\"(?:\\.|[^\"\\])*\")\)"
    )
    for script in soup.find_all("script"):
        raw = script.string or script.get_text()
        chunks = [
            json.loads(match.group(1))
            for match in stream_pattern.finditer(raw)
        ]
        if not chunks:
            continue
        try:
            data = json.loads("".join(chunks))
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(data, list):
            continue

        def resolve(index, depth=0, active=None):
            if (
                depth > 50
                or not isinstance(index, int)
                or index < 0
                or index >= len(data)
            ):
                return index
            if active is None:
                active = set()
            if index in active:
                return None

            node = data[index]
            if isinstance(node, dict):
                active.add(index)
                result = {}
                for key, value in node.items():
                    if not key.startswith("_") or not key[1:].isdigit():
                        continue
                    resolved_key = resolve(int(key[1:]), depth + 1, active)
                    if (
                        isinstance(value, int)
                        and 0 <= value < len(data)
                        and isinstance(data[value], (dict, list, str))
                    ):
                        resolved_value = resolve(value, depth + 1, active)
                    else:
                        resolved_value = value
                    result[str(resolved_key)] = resolved_value
                active.remove(index)
                return result
            if isinstance(node, list):
                active.add(index)
                result = [
                    resolve(value, depth + 1, active)
                    if (
                        isinstance(value, int)
                        and 0 <= value < len(data)
                        and isinstance(data[value], (dict, list, str))
                    )
                    else value
                    for value in node
                ]
                active.remove(index)
                return result
            return node

        records = _assistant_message_records(resolve(0))
        if records:
            candidates.append(records)

    if not candidates:
        return None
    return max(candidates, key=len)[-1]


def _latest_response_record(html):
    soup = BeautifulSoup(html, "html.parser")
    record = _latest_record_from_json_scripts(soup)
    if record:
        return record
    record = _latest_record_from_react_router_stream(soup)
    if record:
        return record

    rendered_messages = soup.select('[data-message-author-role="assistant"]')
    for message in reversed(rendered_messages):
        text = message.get_text("\n", strip=True)
        if text:
            return text, {}

    raise ValueError(
        "No assistant reply was found. Check that the link is a public "
        "ChatGPT shared conversation."
    )


def _image_specs(text, metadata):
    placeholders = list(
        re.finditer(r"<AsyncImage(?:Group)?\b[^>]*?/?>", text)
    )
    components = (
        metadata.get("genui_components", [])
        if isinstance(metadata, dict)
        else []
    )
    components_by_type = {
        component_type: [
            component
            for component in components
            if isinstance(component, dict)
            and component.get("type") == component_type
        ]
        for component_type in ("async_image", "async_image_group")
    }
    placeholders_by_type = {
        component_type: [
            placeholder
            for placeholder in placeholders
            if placeholder.group().startswith("<AsyncImageGroup")
            == (component_type == "async_image_group")
        ]
        for component_type in components_by_type
    }
    if any(
        len(placeholders_by_type[component_type])
        != len(components_by_type[component_type])
        for component_type in components_by_type
    ):
        raise ValueError(
            "The shared reply contains images, but their image data could "
            "not be matched to the reply. Try sharing the conversation again."
        )

    specs = []
    component_indices = {component_type: 0 for component_type in components_by_type}
    for placeholder in placeholders:
        component_type = (
            "async_image_group"
            if placeholder.group().startswith("<AsyncImageGroup")
            else "async_image"
        )
        component_index = component_indices[component_type]
        component_indices[component_type] += 1
        component = components_by_type[component_type][component_index]
        safe_urls = component.get("safe_urls", [])
        image_urls = [
            item.get("url")
            for item in safe_urls
            if isinstance(item, dict)
            and isinstance(item.get("url"), str)
            and item["url"].startswith("https://images.openai.com/")
        ]
        if component_type == "async_image_group":
            group_urls = [
                url for url in image_urls if "purpose=fullsize" in url
            ] or [
                url for url in image_urls if "purpose=inline" in url
            ]
            if not group_urls:
                raise ValueError(
                    "The shared reply contains an image group without "
                    "downloadable ChatGPT images."
                )
            specs.extend((placeholder.group(), url) for url in group_urls)
            continue
        url = next(
            (item for item in image_urls if "purpose=fullsize" in item),
            next(
                (item for item in image_urls if "purpose=inline" in item),
                None,
            ),
        )
        if not url:
            raise ValueError(
                "The shared reply contains an image that ChatGPT did not "
                "provide as a downloadable image."
            )
        specs.append((placeholder.group(), url))
    return specs


def _download_shared_image(placeholder, url):
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "images.openai.com":
        raise ValueError("The shared image URL is not a trusted ChatGPT image.")
    with requests.get(
        url,
        headers={"User-Agent": USER_AGENT},
        timeout=30,
        stream=True,
        allow_redirects=False,
    ) as response:
        response.raise_for_status()
        if response.headers.get("Content-Type", "").split(";")[0].lower() not in {
            "image/jpeg",
            "image/png",
            "image/gif",
            "image/tiff",
            "image/bmp",
        }:
            raise ValueError("ChatGPT returned an unsupported image format.")
        data = bytearray()
        for chunk in response.iter_content(chunk_size=64 * 1024):
            data.extend(chunk)
            if len(data) > MAX_IMAGE_BYTES:
                raise ValueError("A shared image exceeds the 20 MB size limit.")
    return SharedImage(placeholder=placeholder, data=bytes(data))


def _fetch_response_record(url):
    share_url = validate_share_url(url)
    response = requests.get(
        share_url,
        headers={"User-Agent": USER_AGENT},
        timeout=30,
    )
    response.raise_for_status()
    return _latest_response_record(response.text)


def fetch_latest_response_with_images(url):
    text, metadata = _fetch_response_record(url)
    images = [
        _download_shared_image(placeholder, image_url)
        for placeholder, image_url in _image_specs(text, metadata)
    ]
    return text, images


def fetch_latest_response(url):
    text, _metadata = _fetch_response_record(url)
    return text


def extract_latest_response(html):
    text, _metadata = _latest_response_record(html)
    return text


def _clean_inline_text(text):
    text = re.sub(r"<AsyncImage(?:Group)?\b[^>]*?/?>", "", text)
    text = re.sub(
        r"</?(?:box|row|divider)\b[^>]*?/?>",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(
        "\ue200url\ue202(.*?)\ue202.*?\ue201",
        r"\1",
        text,
    )
    text = re.sub(
        "\ue200entity\ue202\\[[^,\\]]+,\\s*\"([^\"]+)\"\\]\ue201",
        r"\1",
        text,
    )
    text = re.sub("\ue200cite\ue202.*?\ue201", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\ue200.*?\ue201", "", text)
    return text


def _main_heading(text):
    in_code_block = False
    has_classification_path = any(
        re.match(
            r"^#{1,3}\s+Recommended Classification Path\s*$",
            line.strip(),
            re.IGNORECASE,
        )
        for line in text.splitlines()
    )
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            continue
        heading = re.match(r"^#\s+(.+)$", line.strip())
        if not heading:
            continue
        title = re.sub(
            r"^\d+(?:\.\d+)*[.)]?\s+",
            "",
            heading.group(1),
        )
        if has_classification_path:
            title_parts = re.split(r"\s+—\s+", title)
            if len(title_parts) >= 3:
                title = title_parts[-1]
        title = _clean_inline_text(title)
        title = re.sub(r"[*_`]", "", title).strip()
        return title
    return ""


def default_output_path(text, desktop=None):
    desktop_path = Path(desktop) if desktop is not None else Path.home() / "Desktop"
    title = _main_heading(text)
    title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", title)
    title = re.sub(r"\s+", " ", title).strip(" .")
    if not title:
        title = "chatgpt-response"
    if re.match(
        r"^(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)",
        title,
        re.IGNORECASE,
    ):
        title = f"_{title}"
    title = title[:120].rstrip(" .") or "chatgpt-response"

    candidate = desktop_path / f"{title}.docx"
    suffix = 2
    while candidate.exists():
        candidate = desktop_path / f"{title} ({suffix}).docx"
        suffix += 1
    return candidate


def _math_element(tag):
    return OxmlElement(f"m:{tag}")


def _math_run(text):
    run = _math_element("r")
    text_element = _math_element("t")
    text_element.set(
        "{http://www.w3.org/XML/1998/namespace}space",
        "preserve",
    )
    text_element.text = text
    run.append(text_element)
    return run


def _append_script_content(container, argument):
    math_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    if argument.tag in {
        f"{{{math_namespace}}}e",
        f"{{{math_namespace}}}sub",
        f"{{{math_namespace}}}sup",
    }:
        for child in list(argument):
            _append_script_content(container, child)
    else:
        container.append(argument)


class _MathParser:
    COMMANDS = {
        "alpha": "α",
        "beta": "β",
        "gamma": "γ",
        "delta": "δ",
        "theta": "θ",
        "lambda": "λ",
        "mu": "μ",
        "pi": "π",
        "rho": "ρ",
        "sigma": "σ",
        "phi": "φ",
        "omega": "ω",
        "cdot": "·",
        "times": "×",
        "pm": "±",
        "quad": "\u2003",
        "qquad": "\u2003\u2003",
        "le": "≤",
        "leq": "≤",
        "ge": "≥",
        "geq": "≥",
        "parallel": "∥",
        "neq": "≠",
        "in": "∈",
        "to": "→",
        "rightarrow": "→",
        "leftarrow": "←",
        "infty": "∞",
        "approx": "≈",
        "circ": "∘",
        "vert": "|",
        "lvert": "|",
        "rvert": "|",
        "{": "{",
        "}": "}",
        "%": "%",
        "_": "_",
        "#": "#",
        "$": "$",
        "&": "&",
    }

    def __init__(self, expression):
        self.tokens = re.findall(
            r"\\[A-Za-z]+|\\.|[A-Za-z]+|\d+(?:\.\d+)?|"
            r"\s+|.",
            expression,
        )
        self.position = 0

    def _group(self, preserve_whitespace=False):
        if self.position < len(self.tokens) and self.tokens[self.position] == "{":
            self.position += 1
            nodes = self._sequence("}", preserve_whitespace)
            if self.position < len(self.tokens):
                self.position += 1
            container = _math_element("e")
            for node in nodes:
                container.append(node)
            return container
        return self._atom()

    def _atom(self):
        if self.position >= len(self.tokens):
            return _math_run("")
        token = self.tokens[self.position]
        self.position += 1
        if token == "{":
            self.position -= 1
            return self._group()

        if token.startswith("\\"):
            command = token[1:]
            if command in {"left", "right"}:
                if self.position < len(self.tokens):
                    delimiter = self.tokens[self.position]
                    self.position += 1
                    if delimiter in {".", "\\."}:
                        return _math_run("")
                    if delimiter.startswith("\\"):
                        delimiter = self.COMMANDS.get(
                            delimiter[1:],
                            delimiter[1:],
                        )
                    return _math_run(delimiter)
                return _math_run("")
            if command == "frac":
                numerator = self._group()
                denominator = self._group()
                fraction = _math_element("f")
                fraction.append(_math_element("fPr"))
                for tag, value in (("num", numerator), ("den", denominator)):
                    part = _math_element(tag)
                    _append_script_content(part, value)
                    fraction.append(part)
                return fraction
            if command == "sqrt":
                radicand = self._group()
                radical = _math_element("rad")
                properties = _math_element("radPr")
                degree_hidden = _math_element("degHide")
                degree_hidden.set(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}val",
                    "1",
                )
                properties.append(degree_hidden)
                radical.append(properties)
                radical.append(_math_element("deg"))
                expression = _math_element("e")
                _append_script_content(expression, radicand)
                radical.append(expression)
                return radical
            if command == "boxed":
                content = self._group()
                border_box = _math_element("borderBox")
                properties = _math_element("borderBoxPr")
                properties.append(_math_element("ctrlPr"))
                border_box.append(properties)
                expression = _math_element("e")
                _append_script_content(expression, content)
                border_box.append(expression)
                return border_box
            if command == "mathbf":
                content = self._group()
                math_namespace = (
                    "http://schemas.openxmlformats.org/officeDocument/2006/math"
                )
                for run in content.iter(f"{{{math_namespace}}}r"):
                    properties = run.find(f"{{{math_namespace}}}rPr")
                    if properties is None:
                        properties = _math_element("rPr")
                        run.insert(0, properties)
                    style = properties.find(f"{{{math_namespace}}}sty")
                    if style is None:
                        style = _math_element("sty")
                        properties.append(style)
                    style.set(
                        f"{{{math_namespace}}}val",
                        "b",
                    )
                return content
            if command in {"text", "mathrm"}:
                return self._group(preserve_whitespace=command == "text")
            if command in self.COMMANDS:
                return _math_run(self.COMMANDS[command])
            return _math_run(command)
        return _math_run(token)

    def _sequence(self, closing=None, preserve_whitespace=False):
        nodes = []
        while self.position < len(self.tokens):
            token = self.tokens[self.position]
            if token == closing:
                break
            if token.isspace() and not preserve_whitespace:
                self.position += 1
                continue
            if token in {"^", "_"} and nodes:
                self.position += 1
                script = self._group()
                base = nodes.pop()
                has_sub = base.tag.endswith(("}sSub", "}sSubSup"))
                has_sup = base.tag.endswith(("}sSup", "}sSubSup"))
                if has_sub or has_sup:
                    namespaces = {
                        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"
                    }
                    old_expression = base.find("m:e", namespaces=namespaces)
                    old_sub = base.find("m:sub", namespaces=namespaces)
                    old_sup = base.find("m:sup", namespaces=namespaces)
                    needs_both_scripts = (
                        token == "^" and has_sub
                    ) or (
                        token == "_" and has_sup
                    )
                    if needs_both_scripts:
                        script_tag = "sSubSup"
                    else:
                        script_tag = "sSup" if token == "^" else "sSub"
                    script_node = _math_element(script_tag)
                    if script_node.tag.endswith("sSubSup"):
                        script_node.append(_math_element("sSubSupPr"))
                    script_node.append(old_expression)
                    sub = script if token == "_" else old_sub
                    sup = script if token == "^" else old_sup
                    if sub is not None:
                        part = _math_element("sub")
                        _append_script_content(part, sub)
                        script_node.append(part)
                    if sup is not None:
                        part = _math_element("sup")
                        _append_script_content(part, sup)
                        script_node.append(part)
                    base = script_node
                else:
                    script_node = _math_element("sSup" if token == "^" else "sSub")
                    script_node.append(_math_element("sSupPr" if token == "^" else "sSubPr"))
                    expression = _math_element("e")
                    _append_script_content(expression, base)
                    script_node.append(expression)
                    part = _math_element("sup" if token == "^" else "sub")
                    _append_script_content(part, script)
                    script_node.append(part)
                    base = script_node
                nodes.append(base)
                continue
            nodes.append(self._atom())
        return nodes

    def parse(self):
        expression = _math_element("oMath")
        for node in self._sequence():
            expression.append(node)
        return expression


def _add_math(paragraph, expression, display=False, alignment=None):
    math = _MathParser(expression).parse()
    if display:
        display_math = _math_element("oMathPara")
        if alignment is not None:
            properties = _math_element("oMathParaPr")
            justification = _math_element("jc")
            justification.set(
                "{http://schemas.openxmlformats.org/officeDocument/2006/math}val",
                alignment,
            )
            properties.append(justification)
            display_math.append(properties)
        display_math.append(math)
        paragraph._p.append(display_math)
    else:
        paragraph._p.append(math)


def _add_pasted_text(document, paragraphs):
    for pasted_paragraph in paragraphs:
        paragraph = document.add_paragraph()
        for pasted_run in pasted_paragraph.runs:
            run = paragraph.add_run(pasted_run.text)
            formatting = pasted_run.format
            run.bold = formatting.bold
            run.italic = formatting.italic
            run.underline = formatting.underline
            if formatting.color:
                run.font.color.rgb = RGBColor.from_string(formatting.color)
            if formatting.font_name:
                run.font.name = formatting.font_name
            if formatting.font_size:
                run.font.size = Pt(formatting.font_size)
            if formatting.background:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), formatting.background)
                run._r.get_or_add_rPr().append(shading)


def _numbered_section_heading(text):
    stripped = re.sub(r"^#{1,6}\s+", "", text.strip())
    return re.fullmatch(
        r"(?:\*\*)?(\d+(?:\.\d+)+\.?)\s+(.+?)(?:\*\*)?",
        stripped,
    )


def _numbered_heading_level(section_heading):
    if section_heading is None:
        return None
    return len(section_heading.group(1).rstrip(".").split(".")) + 1


def _heading_level(line):
    stripped = line.strip()
    heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
    if heading:
        numbered_heading = _numbered_section_heading(stripped)
        return _numbered_heading_level(numbered_heading) or len(
            heading.group(1)
        )
    return _numbered_heading_level(_numbered_section_heading(stripped))


def _next_content_is_subheading(lines, start_index, heading_level):
    for line in lines[start_index + 1 :]:
        stripped = line.strip()
        if (
            not stripped
            or re.match(r"^(---+|___+|\*\*\*+)\s*$", stripped)
            or re.fullmatch(
                r"</?(?:box|row|divider)\b[^>]*?/?>",
                stripped,
                re.IGNORECASE,
            )
        ):
            continue
        if stripped.startswith("```"):
            return False
        return _heading_level(line) == heading_level + 1
    return False


def _strip_redundant_table_symbol_parentheses(expression):
    match = re.fullmatch(
        r"\(\s*([A-Za-z]_\{?[A-Za-z0-9]+\}?)\s*\)",
        expression.strip(),
    )
    return match.group(1) if match else expression


def _add_inline_runs(paragraph, text, table_math=False):
    text = text.strip()
    token_pattern = re.compile(
        r"(\\\(.+?\\\)|\\\[.+?\\\]|\$.+?\$|\*\*.+?\*\*|__.+?__|"
        r"(?<!\*)\*[^*]+\*(?!\*)|(?<!_)_[^_]+_(?!_)|`[^`]+`)"
    )
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(_clean_inline_text(text[position:match.start()]))
        token = match.group()
        if token.startswith((r"\(", r"\[")):
            expression = token[2:-2]
            if table_math:
                expression = _strip_redundant_table_symbol_parentheses(
                    expression
                )
            _add_math(
                paragraph,
                expression,
                display=table_math,
                alignment="left" if table_math else None,
            )
        elif token.startswith("$"):
            expression = token[1:-1]
            if table_math:
                expression = _strip_redundant_table_symbol_parentheses(
                    expression
                )
            _add_math(
                paragraph,
                expression,
                display=table_math,
                alignment="left" if table_math else None,
            )
        elif token.startswith(("**", "__")):
            run = paragraph.add_run(_clean_inline_text(token[2:-2]))
            run.bold = True
        elif token.startswith(("`",)):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(_clean_inline_text(text[position:]))
    while paragraph.runs and paragraph._p[-1] is paragraph.runs[-1]._r:
        run = paragraph.runs[-1]
        trimmed = run.text.rstrip()
        if trimmed == run.text:
            break
        if trimmed:
            run.text = trimmed
            break
        run._r.getparent().remove(run._r)


def _is_bare_math_table_cell(text):
    value = text.strip()
    if not value:
        return False
    if re.search(r"(?:\*\*.+?\*\*|__.+__)", value):
        return False
    if re.fullmatch(r"(?:\\\(.+?\\\)|\\\[.+?\\\]|\$.+?\$)", value):
        return False
    if re.fullmatch(r"(?:[A-Za-z](?:_\{?[A-Za-z0-9]+\}?)?|\d+[A-Za-z])", value):
        return True
    has_latex_syntax = re.search(
        r"\\[A-Za-z]+|[_^{}=+*/<>≤≥]|\s-\s",
        value,
    )
    return bool(
        has_latex_syntax
        and re.fullmatch(
            r"[A-Za-z0-9\s\\{}_^.,=+\-*/()|<>≤≥πρ]+",
            value,
        )
    )


def _add_response_text(document, text, images=(), heading_content=()):
    images_by_placeholder = {}
    for image in images:
        images_by_placeholder.setdefault(image.placeholder, []).append(image)
    in_code_block = False
    main_heading_seen = False
    main_heading_image_window = False
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if re.match(r"^If you want, I can:\s*$", line.strip(), re.IGNORECASE):
            suggestions = [
                candidate.strip()
                for candidate in lines[index + 1 :]
                if candidate.strip()
            ]
            if suggestions and all(
                re.match(r"^[-*+]\s+", suggestion)
                for suggestion in suggestions
            ):
                lines = lines[:index]
                break

    has_classification_path = any(
        re.match(
            r"^#{1,3}\s+Recommended Classification Path\s*$",
            line.strip(),
            re.IGNORECASE,
        )
        for line in lines
    )
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            index += 1
            continue
        placeholder_images = images_by_placeholder.get(stripped)
        if placeholder_images:
            paragraph = (
                document.paragraphs[-1]
                if main_heading_image_window
                and document.paragraphs
                and re.match(
                    r"^#?\s*Source:\s*AI Source:",
                    document.paragraphs[-1].text.strip(),
                    re.IGNORECASE,
                )
                else document.add_paragraph()
            )
            for image in placeholder_images:
                paragraph.add_run().add_picture(
                    io.BytesIO(image.data),
                    width=Inches(6),
                )
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if re.fullmatch(r"</?(?:box|row|divider)\b[^>]*?/?>", stripped, re.IGNORECASE):
            index += 1
            continue
        if in_code_block:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            index += 1
            continue

        source_metadata_line = re.match(
            r"^#?\s*Source:\s*AI Source:",
            stripped,
            re.IGNORECASE,
        )
        if main_heading_image_window and not source_metadata_line:
            main_heading_image_window = False

        if re.match(r"^(---+|___+|\*\*\*+)\s*$", stripped):
            index += 1
            continue
        if stripped in {r"\[", "$$"}:
            closing = r"\]" if stripped == r"\[" else "$$"
            equation_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != closing:
                equation_lines.append(lines[index].strip())
                index += 1
            if index < len(lines):
                index += 1
            equation = " ".join(equation_lines)
            paragraph = document.add_paragraph()
            if equation:
                _add_math(paragraph, equation, display=True)
            continue

        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            if re.match(r"^\|?[\s:|-]+\|?$", separator) and "-" in separator:
                rows = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    cells = [
                        cell.strip()
                        for cell in lines[index].strip().strip("|").split("|")
                    ]
                    first_cell = lines[index].strip().strip("|").split("|")[0]
                    if not re.match(r"^\s*:?-+:?\s*$", first_cell):
                        rows.append(cells)
                    index += 1
                if len(rows) > 1:
                    table = document.add_table(
                        rows=len(rows),
                        cols=max(len(row) for row in rows),
                    )
                    table.style = "Normal Table"
                    table.alignment = WD_TABLE_ALIGNMENT.LEFT
                    if len(table.columns) == 4:
                        column_widths = (0.85, 2.31, 0.64, 0.75)
                        table.autofit = False
                        for column, width in zip(table.columns, column_widths):
                            column.width = Inches(width)
                        for row in table.rows:
                            for cell, width in zip(row.cells, column_widths):
                                cell.width = Inches(width)
                    for row_index, cells in enumerate(rows):
                        for column_index, value in enumerate(cells):
                            if row_index == 1 and re.match(r"^\s*:?-+:?\s*$", value):
                                continue
                            value = re.sub(
                                r"^\*\*Basediameter\*\*\(\(([A-Za-z])\)\)$",
                                r"**Base diameter** (\\(\1\\))",
                                value,
                            )
                            cell = table.cell(row_index, column_index)
                            paragraph = cell.paragraphs[0]
                            if row_index > 0 and _is_bare_math_table_cell(value):
                                value = (
                                    _strip_redundant_table_symbol_parentheses(
                                        value
                                    )
                                )
                                _add_math(
                                    paragraph,
                                    value,
                                    display=True,
                                    alignment="left",
                                )
                            else:
                                _add_inline_runs(
                                    paragraph,
                                    value,
                                    table_math=True,
                                )
                            if row_index == 0:
                                for run in paragraph.runs:
                                    run.bold = True
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    table.alignment = WD_TABLE_ALIGNMENT.LEFT
                    continue
                continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        numbered_section_heading = _numbered_section_heading(stripped)
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if heading:
            if len(heading.group(1)) == 1 and not main_heading_seen:
                main_heading_seen = True
                main_heading_image_window = True
            heading_level = (
                _numbered_heading_level(numbered_section_heading)
                or len(heading.group(1))
            )
            paragraph = document.add_heading(level=heading_level)
            if numbered_section_heading:
                heading_text = numbered_section_heading.group(2)
            else:
                heading_text = re.sub(
                    r"^\d+(?:\.\d+)*[.)]?\s+",
                    "",
                    heading.group(2),
                )
            if len(heading.group(1)) == 1 and has_classification_path:
                title_parts = re.split(r"\s+—\s+", heading_text)
                if len(title_parts) >= 3:
                    heading_text = title_parts[-1]
            _add_inline_runs(
                paragraph,
                heading_text,
            )
            if not _next_content_is_subheading(
                lines,
                index,
                heading_level,
            ):
                _add_pasted_text(document, heading_content)
        elif numbered_section_heading:
            heading_level = _numbered_heading_level(
                numbered_section_heading
            )
            paragraph = document.add_heading(level=heading_level)
            _add_inline_runs(paragraph, numbered_section_heading.group(2))
            if not _next_content_is_subheading(
                lines,
                index,
                heading_level,
            ):
                _add_pasted_text(document, heading_content)
        elif bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, numbered.group(1))
        else:
            if stripped.startswith(r"\]") or stripped == "$$":
                index += 1
                continue
            inline_display = re.fullmatch(r"\\\[(.+?)\\\]", stripped)
            paragraph = document.add_paragraph()
            if inline_display:
                _add_math(paragraph, inline_display.group(1), display=True)
            else:
                _add_inline_runs(paragraph, line)
        index += 1


def _apply_document_format(document):
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Roboto Mono"
    normal.font.size = Pt(12)

    heading_styles = (
        ("Heading 1", 20, 18),
        ("Heading 2", 16, 8),
        ("Heading 3", 14, 8),
        ("Heading 4", 13, 6),
        ("Heading 5", 12, 4),
        ("Heading 6", 11, 4),
    )
    for name, size, space_before in heading_styles:
        style = document.styles[name]
        style.font.name = "Roboto Mono"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def save_docx(text, destination, images=(), heading_content=()):
    document = Document()
    _add_response_text(document, text, images, heading_content)
    _apply_document_format(document)
    compatibility = document.settings._element.find(qn("w:compat"))
    compatibility_mode = next(
        (
            setting
            for setting in compatibility.findall(qn("w:compatSetting"))
            if setting.get(qn("w:name")) == "compatibilityMode"
        ),
        None,
    )
    if compatibility_mode is None:
        compatibility_mode = OxmlElement("w:compatSetting")
        compatibility_mode.set(qn("w:name"), "compatibilityMode")
        compatibility_mode.set(
            qn("w:uri"),
            "http://schemas.microsoft.com/office/word",
        )
        compatibility.append(compatibility_mode)
    compatibility_mode.set(qn("w:val"), "15")

    package_buffer = io.BytesIO()
    document.save(package_buffer)
    package_buffer.seek(0)
    finalized_package = io.BytesIO()
    with ZipFile(package_buffer) as package:
        with ZipFile(finalized_package, "w", ZIP_DEFLATED) as finalized:
            for entry in package.infolist():
                data = package.read(entry.filename)
                if entry.filename == "docProps/app.xml":
                    properties = etree.fromstring(data)
                    namespace = (
                        "{http://schemas.openxmlformats.org/officeDocument/"
                        "2006/extended-properties}"
                    )
                    application = properties.find(f"{namespace}Application")
                    version = properties.find(f"{namespace}AppVersion")
                    if application is None or version is None:
                        raise ValueError(
                            "The Word document application metadata is missing."
                        )
                    application.text = "Microsoft Office Word"
                    version.text = "16.0000"
                    data = etree.tostring(
                        properties,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                finalized.writestr(entry, data)
    with open(destination, "wb") as output:
        output.write(finalized_package.getvalue())


class ChatGPTToDocxApp:
    def __init__(self, root):
        self.root = root
        root.title("ChatGPT to Word")
        root.minsize(680, 500)

        frame = ttk.Frame(root, padding=16)
        frame.grid(sticky="nsew")
        root.columnconfigure(0, weight=1)
        root.rowconfigure(0, weight=1)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(3, weight=1)

        ttk.Label(frame, text="ChatGPT share link:").grid(
            row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 12)
        )
        self.url_entry = ttk.Entry(frame, width=58)
        self.url_entry.grid(row=0, column=1, sticky="ew", pady=(0, 12))

        ttk.Label(frame, text="Save Word file to:").grid(
            row=1, column=0, sticky="w", padx=(0, 10), pady=(0, 12)
        )
        self.destination_entry = ttk.Entry(frame)
        self.destination_entry.grid(row=1, column=1, sticky="ew", pady=(0, 12))
        ttk.Button(frame, text="Browse...", command=self.choose_destination).grid(
            row=1, column=2, padx=(8, 0), pady=(0, 12)
        )

        ttk.Label(
            frame,
            text="Formatted text after headings (skip headings with subheadings):",
        ).grid(row=2, column=0, columnspan=2, sticky="w", pady=(8, 4))
        ttk.Button(
            frame,
            text="Clear",
            command=self.clear_pasted_text,
        ).grid(row=2, column=2, sticky="e", pady=(8, 4))
        pasted_frame = ttk.Frame(frame)
        pasted_frame.grid(row=3, column=0, columnspan=3, sticky="nsew")
        pasted_frame.rowconfigure(0, weight=1)
        pasted_frame.columnconfigure(0, weight=1)
        self.pasted_text = tk.Text(
            pasted_frame,
            width=70,
            height=9,
            wrap="word",
            undo=True,
        )
        self.pasted_text.grid(row=0, column=0, sticky="nsew")
        pasted_scrollbar = ttk.Scrollbar(
            pasted_frame,
            orient="vertical",
            command=self.pasted_text.yview,
        )
        pasted_scrollbar.grid(row=0, column=1, sticky="ns")
        self.pasted_text.configure(yscrollcommand=pasted_scrollbar.set)
        self._preview_formats = {}
        self._formats_by_tag = {}
        self._preview_fonts = []
        self.pasted_text.bind("<Control-v>", self.paste_formatted_text)
        self.pasted_text.bind("<Control-V>", self.paste_formatted_text)
        self.pasted_text.bind("<Shift-Insert>", self.paste_formatted_text)
        self.pasted_text.bind("<<Paste>>", self.paste_formatted_text)

        self.convert_button = ttk.Button(
            frame, text="Convert latest reply", command=self.convert
        )
        self.convert_button.grid(row=4, column=1, sticky="e", pady=(10, 0))
        self.progress = ttk.Progressbar(
            frame,
            mode="indeterminate",
            maximum=100,
        )
        self.progress.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(10, 0))
        self.progress.grid_remove()
        self.status = ttk.Label(frame, text="Only public ChatGPT share links are supported.")
        self.status.grid(row=6, column=0, columnspan=3, sticky="w", pady=(12, 0))

    def clear_pasted_text(self):
        self.pasted_text.delete("1.0", "end")
        self.pasted_text.focus_set()

    def _format_tag(self, formatting):
        tag = self._preview_formats.get(formatting)
        if tag is not None:
            return tag
        tag = f"pasted_format_{len(self._preview_formats)}"
        self._preview_formats[formatting] = tag
        self._formats_by_tag[tag] = formatting
        options = {}
        if formatting.color:
            options["foreground"] = f"#{formatting.color}"
        if formatting.background:
            options["background"] = f"#{formatting.background}"
        if (
            formatting.font_name
            or formatting.font_size
            or formatting.bold
            or formatting.italic
            or formatting.underline
        ):
            font = tkfont.Font(
                self.pasted_text,
                family=formatting.font_name or "Segoe UI",
                size=round(formatting.font_size or 10),
                weight="bold" if formatting.bold else "normal",
                slant="italic" if formatting.italic else "roman",
                underline=formatting.underline,
            )
            self._preview_fonts.append(font)
            options["font"] = font
        if options:
            self.pasted_text.tag_configure(tag, **options)
        return tag

    def _insert_pasted_paragraphs(self, paragraphs):
        if not paragraphs:
            return
        selection = self.pasted_text.tag_ranges("sel")
        if selection:
            self.pasted_text.delete(selection[0], selection[1])
        for paragraph_index, paragraph in enumerate(paragraphs):
            for run in paragraph.runs:
                tag = self._format_tag(run.format)
                self.pasted_text.insert("insert", run.text, (tag,))
            if paragraph_index < len(paragraphs) - 1:
                self.pasted_text.insert("insert", "\n")

    def paste_formatted_text(self, _event=None):
        html = _clipboard_html()
        paragraphs = parse_pasted_html(html) if html else ()
        if not paragraphs:
            try:
                plain_text = self.root.clipboard_get()
            except tk.TclError:
                return "break"
            paragraphs = tuple(
                PastedTextParagraph((PastedTextRun(line),))
                for line in plain_text.splitlines()
            )
        self._insert_pasted_paragraphs(paragraphs)
        return "break"

    def _pasted_text_content(self):
        paragraphs = []
        line_count = int(self.pasted_text.index("end-1c").split(".")[0])
        for line_number in range(1, line_count + 1):
            line = self.pasted_text.get(
                f"{line_number}.0",
                f"{line_number}.end",
            )
            runs = []
            for column, character in enumerate(line):
                tag = next(
                    (
                        tag_name
                        for tag_name in self.pasted_text.tag_names(
                            f"{line_number}.{column}"
                        )
                        if tag_name in self._formats_by_tag
                    ),
                    None,
                )
                formatting = (
                    self._formats_by_tag[tag]
                    if tag
                    else PastedTextFormat()
                )
                if runs and runs[-1].format == formatting:
                    previous = runs[-1]
                    runs[-1] = PastedTextRun(
                        previous.text + character,
                        formatting,
                    )
                else:
                    runs.append(PastedTextRun(character, formatting))
            paragraphs.append(PastedTextParagraph(tuple(runs)))
        while paragraphs and not any(
            run.text for run in paragraphs[-1].runs
        ):
            paragraphs.pop()
        return tuple(paragraphs)

    def choose_destination(self):
        destination = filedialog.asksaveasfilename(
            title="Save Word document",
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
            initialfile="chatgpt-response.docx",
        )
        if destination:
            self.destination_entry.delete(0, tk.END)
            self.destination_entry.insert(0, destination)

    def show_error_dialog(self, error):
        error_text = f"{type(error).__name__}: {error}"
        dialog = tk.Toplevel(self.root)
        dialog.title("Could not convert link")
        dialog.transient(self.root)
        dialog.resizable(False, False)

        frame = ttk.Frame(dialog, padding=16)
        frame.grid(sticky="nsew")
        ttk.Label(
            frame,
            text="The conversion failed. Copy the error below if you need help:",
            wraplength=420,
        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        details = tk.Text(frame, width=58, height=5, wrap="word")
        details.insert("1.0", error_text)
        details.configure(state="disabled")
        details.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 12))

        def copy_error():
            self.root.clipboard_clear()
            self.root.clipboard_append(error_text)
            self.root.update()
            dialog.destroy()
            self.status.configure(text="Error copied to clipboard.")

        ttk.Button(frame, text="Copy error", command=copy_error).grid(
            row=2, column=0, sticky="e", padx=(0, 8)
        )
        ttk.Button(frame, text="Close", command=dialog.destroy).grid(
            row=2, column=1, sticky="e"
        )
        dialog.bind("<Escape>", lambda _event: dialog.destroy())
        dialog.bind("<Return>", lambda _event: dialog.destroy())
        dialog.grab_set()
        dialog.focus_set()

    def convert(self):
        url = self.url_entry.get().strip()
        destination = self.destination_entry.get().strip()
        heading_content = self._pasted_text_content()

        self.convert_button.configure(state="disabled")
        self.status.configure(text="Fetching the latest reply...")
        self.progress.grid()
        self.progress.start(12)
        worker = threading.Thread(
            target=self._convert_in_background,
            args=(url, destination, heading_content),
            daemon=True,
        )
        worker.start()

    def _convert_in_background(self, url, destination, heading_content):
        try:
            text, images = fetch_latest_response_with_images(url)
            if not destination:
                destination = str(default_output_path(text))
            save_docx(
                text,
                destination,
                images,
                heading_content,
            )
        except (ValueError, requests.RequestException, OSError) as error:
            self.root.after(
                0,
                self._finish_conversion,
                None,
                error,
            )
        else:
            self.root.after(
                0,
                self._finish_conversion,
                destination,
                None,
            )

    def _finish_conversion(self, destination, error):
        self.progress.stop()
        self.progress.grid_remove()
        self.convert_button.configure(state="normal")
        if error is not None:
            self.status.configure(text="Conversion failed.")
            self.show_error_dialog(error)
            return

        if not self.destination_entry.get().strip():
            self.destination_entry.insert(0, destination)
        self.status.configure(text=f"Saved: {destination}")
        messagebox.showinfo(
            "Conversion complete",
            f"Word document saved to:\n{destination}",
        )


def main():
    root = tk.Tk()
    ChatGPTToDocxApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
